#!/usr/bin/env python3
"""Convierte la documentacion Markdown del taller a un archivo .docx.

Uso:
    python scripts/md_a_docx.py docs/documentacion.md docs/documentacion.docx

Soporta lo que usa la documentacion: encabezados, parrafos con negrita/cursiva/
codigo/enlaces, listas con y sin numeracion, citas, tablas de tuberias, reglas
horizontales y bloques de codigo cercados (incluidos los diagramas ASCII, que se
componen en fuente monoespaciada pequena para que no se rompa el trazado).

Requiere: pip install python-docx
"""
import re
import sys

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

FUENTE_CODIGO = "Consolas"
GRIS_CODIGO = "F2F2F2"


def sombrear(parrafo, color_hex):
    """Aplica un color de fondo al parrafo (bloques de codigo)."""
    sombra = OxmlElement("w:shd")
    sombra.set(qn("w:val"), "clear")
    sombra.set(qn("w:fill"), color_hex)
    parrafo._p.get_or_add_pPr().append(sombra)


def agregar_borde_inferior(parrafo):
    """Dibuja una linea horizontal (equivalente de --- en Markdown)."""
    bordes = OxmlElement("w:pBdr")
    inferior = OxmlElement("w:bottom")
    inferior.set(qn("w:val"), "single")
    inferior.set(qn("w:sz"), "6")
    inferior.set(qn("w:color"), "BFBFBF")
    bordes.append(inferior)
    parrafo._p.get_or_add_pPr().append(bordes)


# Negrita, cursiva, codigo y enlaces en linea.
PATRON_INLINE = re.compile(
    r"(\*\*.+?\*\*|__.+?__|`[^`]+`|\*[^*\n]+\*|_[^_\n]+_|\[[^\]]+\]\([^)]+\))"
)


def escribir_texto(parrafo, texto, negrita_base=False):
    """Agrega texto interpretando el formato en linea de Markdown."""
    for trozo in PATRON_INLINE.split(texto):
        if not trozo:
            continue
        corrida = None
        if trozo.startswith("**") and trozo.endswith("**") and len(trozo) > 4:
            corrida = parrafo.add_run(trozo[2:-2])
            corrida.bold = True
        elif trozo.startswith("__") and trozo.endswith("__") and len(trozo) > 4:
            corrida = parrafo.add_run(trozo[2:-2])
            corrida.bold = True
        elif trozo.startswith("`") and trozo.endswith("`") and len(trozo) > 2:
            corrida = parrafo.add_run(trozo[1:-1])
            corrida.font.name = FUENTE_CODIGO
            corrida.font.size = Pt(9)
            corrida.font.color.rgb = RGBColor(0xC0, 0x39, 0x2B)
        elif trozo.startswith("[") and "](" in trozo:
            etiqueta = trozo[1 : trozo.index("](")]
            corrida = parrafo.add_run(etiqueta)
            corrida.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)
            corrida.underline = True
        elif len(trozo) > 2 and trozo[0] in "*_" and trozo[-1] == trozo[0]:
            corrida = parrafo.add_run(trozo[1:-1])
            corrida.italic = True
        else:
            corrida = parrafo.add_run(trozo)
        if negrita_base and corrida is not None and corrida.bold is None:
            corrida.bold = True


def es_separador_de_tabla(linea):
    return bool(re.match(r"^\s*\|[\s:\-|]+\|\s*$", linea))


def celdas(linea):
    return [c.strip() for c in linea.strip().strip("|").split("|")]


def agregar_tabla(documento, filas):
    """Convierte una tabla de tuberias de Markdown en una tabla de Word."""
    encabezado = celdas(filas[0])
    cuerpo = [celdas(f) for f in filas[2:]]
    tabla = documento.add_table(rows=1, cols=len(encabezado))
    tabla.style = "Light Grid Accent 1"
    tabla.alignment = WD_TABLE_ALIGNMENT.CENTER

    for i, titulo in enumerate(encabezado):
        celda = tabla.rows[0].cells[i]
        celda.text = ""
        escribir_texto(celda.paragraphs[0], titulo, negrita_base=True)
        for corrida in celda.paragraphs[0].runs:
            corrida.bold = True
            corrida.font.size = Pt(9)

    for fila in cuerpo:
        nueva = tabla.add_row().cells
        for i, valor in enumerate(fila[: len(encabezado)]):
            nueva[i].text = ""
            escribir_texto(nueva[i].paragraphs[0], valor)
            for corrida in nueva[i].paragraphs[0].runs:
                corrida.font.size = Pt(9)
    documento.add_paragraph()


def agregar_codigo(documento, lineas):
    """Inserta un bloque de codigo o un diagrama ASCII."""
    ancho_max = max((len(l) for l in lineas), default=0)
    tamano = 8 if ancho_max <= 80 else 6.5
    for linea in lineas:
        parrafo = documento.add_paragraph()
        parrafo.paragraph_format.space_after = Pt(0)
        parrafo.paragraph_format.space_before = Pt(0)
        parrafo.paragraph_format.left_indent = Inches(0.15)
        corrida = parrafo.add_run(linea if linea else " ")
        corrida.font.name = FUENTE_CODIGO
        corrida.font.size = Pt(tamano)
        sombrear(parrafo, GRIS_CODIGO)
    documento.add_paragraph()


def convertir(ruta_md, ruta_docx):
    lineas = open(ruta_md, encoding="utf-8").read().split("\n")
    documento = Document()

    estilo = documento.styles["Normal"]
    estilo.font.name = "Calibri"
    estilo.font.size = Pt(10.5)
    for seccion in documento.sections:
        seccion.left_margin = seccion.right_margin = Inches(0.8)
        seccion.top_margin = seccion.bottom_margin = Inches(0.8)

    i = 0
    while i < len(lineas):
        linea = lineas[i]

        # Bloques de codigo cercados.
        if linea.startswith("```"):
            j = i + 1
            bloque = []
            while j < len(lineas) and not lineas[j].startswith("```"):
                bloque.append(lineas[j])
                j += 1
            agregar_codigo(documento, bloque)
            i = j + 1
            continue

        # Tablas de tuberias.
        if linea.strip().startswith("|") and i + 1 < len(lineas) and es_separador_de_tabla(lineas[i + 1]):
            j = i
            filas = []
            while j < len(lineas) and lineas[j].strip().startswith("|"):
                filas.append(lineas[j])
                j += 1
            agregar_tabla(documento, filas)
            i = j
            continue

        # Reglas horizontales.
        if re.match(r"^\s*(-{3,}|\*{3,})\s*$", linea):
            agregar_borde_inferior(documento.add_paragraph())
            i += 1
            continue

        # Encabezados.
        encabezado = re.match(r"^(#{1,6})\s+(.*)$", linea)
        if encabezado:
            nivel = len(encabezado.group(1))
            parrafo = documento.add_heading("", level=min(nivel, 4))
            escribir_texto(parrafo, encabezado.group(2))
            i += 1
            continue

        # Citas.
        if linea.startswith("> "):
            parrafo = documento.add_paragraph(style="Intense Quote")
            escribir_texto(parrafo, linea[2:])
            i += 1
            continue

        # Listas con vinetas (admite sangria de continuacion).
        vineta = re.match(r"^(\s*)[-*+]\s+(.*)$", linea)
        if vineta:
            parrafo = documento.add_paragraph(style="List Bullet")
            if len(vineta.group(1)) >= 2:
                parrafo.paragraph_format.left_indent = Inches(0.6)
            escribir_texto(parrafo, vineta.group(2))
            i += 1
            while i < len(lineas) and lineas[i].startswith("  ") and lineas[i].strip() and not re.match(r"^\s*([-*+]|\d+\.)\s", lineas[i]):
                escribir_texto(parrafo, " " + lineas[i].strip())
                i += 1
            continue

        # Listas numeradas.
        numerada = re.match(r"^(\s*)\d+\.\s+(.*)$", linea)
        if numerada:
            parrafo = documento.add_paragraph(style="List Number")
            escribir_texto(parrafo, numerada.group(2))
            i += 1
            while i < len(lineas) and lineas[i].startswith("   ") and lineas[i].strip() and not re.match(r"^\s*([-*+]|\d+\.)\s", lineas[i]):
                escribir_texto(parrafo, " " + lineas[i].strip())
                i += 1
            continue

        # Linea en blanco.
        if not linea.strip():
            i += 1
            continue

        # Parrafo normal (se unen las lineas hasta el proximo elemento).
        bloque = [linea.strip()]
        i += 1
        while i < len(lineas) and lineas[i].strip() and not re.match(
            r"^(#{1,6}\s|```|\||\s*[-*+]\s|\s*\d+\.\s|>\s|-{3,}$)", lineas[i]
        ):
            bloque.append(lineas[i].strip())
            i += 1
        parrafo = documento.add_paragraph()
        parrafo.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        escribir_texto(parrafo, " ".join(bloque))

    documento.save(ruta_docx)
    print(f"generado {ruta_docx}")


if __name__ == "__main__":
    entrada = sys.argv[1] if len(sys.argv) > 1 else "docs/documentacion.md"
    salida = sys.argv[2] if len(sys.argv) > 2 else "docs/documentacion.docx"
    convertir(entrada, salida)
